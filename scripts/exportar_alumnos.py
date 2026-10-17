import json

import openpyxl

wb = openpyxl.load_workbook("docs/ALUMNOS.xlsx", data_only=True)
ws = wb["CUOTA 2026"]


def txt(v):
    if v is None:
        return None
    s = str(v).strip()
    return s if s else None


def fecha(v):
    if v is None:
        return None
    if hasattr(v, "strftime"):
        return v.strftime("%Y-%m-%d")
    s = str(v)[:10]
    return s if len(s) == 10 else None


filas = []
for r in range(3, ws.max_row + 1):
    nombres = txt(ws.cell(row=r, column=4).value)
    if not nombres:
        continue
    filas.append(
        {
            "nombres": nombres,
            "apellidos": txt(ws.cell(row=r, column=5).value),
            "sede": txt(ws.cell(row=r, column=3).value),
            "nacimiento": fecha(ws.cell(row=r, column=6).value),
            "ci": txt(ws.cell(row=r, column=8).value),
            "nacionalidad": txt(ws.cell(row=r, column=9).value),
            "sexo": txt(ws.cell(row=r, column=10).value),
            "turno": txt(ws.cell(row=r, column=11).value),
            "rue": txt(ws.cell(row=r, column=12).value),
            "beca": ws.cell(row=r, column=13).value,
            "presupuesto": ws.cell(row=r, column=14).value,
            "matricula": ws.cell(row=r, column=15).value,
        }
    )

with open("docs/_alumnos.json", "w", encoding="utf-8") as f:
    json.dump(filas, f, ensure_ascii=False, indent=2, default=str)

print("Alumnos exportados:", len(filas))
