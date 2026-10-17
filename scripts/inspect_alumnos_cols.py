from collections import Counter

import openpyxl

wb = openpyxl.load_workbook("docs/ALUMNOS.xlsx", data_only=True)
ws = wb["CUOTA 2026"]

hdr = [ws.cell(row=2, column=c).value for c in range(1, ws.max_column + 1)]
print("Columnas:", len(hdr))
for i, h in enumerate(hdr, 1):
    print(f"  {i:2d}. {h!r}")


def distintos(col):
    c = Counter(
        str(ws.cell(row=r, column=col).value).strip()
        for r in range(3, ws.max_row + 1)
        if ws.cell(row=r, column=col).value not in (None, "")
    )
    return dict(sorted(c.items(), key=lambda kv: -kv[1]))


for col, nombre in [(3, "SEDE"), (9, "NACIONALIDAD"), (10, "SEXO"), (11, "TURNOS"), (12, "RUE"), (13, "BECA")]:
    print(nombre, "->", distintos(col))

cnt = sum(1 for r in range(3, ws.max_row + 1) if ws.cell(row=r, column=4).value not in (None, ""))
print("Alumnos (filas con NOMBRES):", cnt)
