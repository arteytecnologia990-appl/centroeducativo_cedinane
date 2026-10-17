import json
import sys

import openpyxl

path = "docs/COLABORADORES.xlsx"
wb = openpyxl.load_workbook(path, data_only=False)
print("Hojas:", wb.sheetnames)

objetivo = None
for nombre in wb.sheetnames:
    n = nombre.lower()
    if "personal" in n or "datos" in n:
        objetivo = nombre
        break
print("Hoja objetivo:", objetivo)

ws = wb[objetivo]
print("Dimensiones:", ws.dimensions, "| filas:", ws.max_row, "| columnas:", ws.max_column)

mergeadas = [str(r) for r in ws.merged_cells.ranges]
print("Rangos combinados:", mergeadas[:10])

print("\nEncabezados (fila 1):")
headers = [c.value for c in ws[1]]
for i, h in enumerate(headers, 1):
    print(f"  {i:2d}. {h!r}")

print("\nPrimeras filas de datos:")
for r in range(2, min(ws.max_row, 7) + 1):
    vals = [ws.cell(row=r, column=c).value for c in range(1, ws.max_column + 1)]
    print(f"  fila {r}: {vals}")

print("\nUltima fila no vacia:")
for r in range(ws.max_row, 1, -1):
    vals = [ws.cell(row=r, column=c).value for c in range(1, ws.max_column + 1)]
    if any(v not in (None, "") for v in vals):
        print(f"  fila {r}: {vals}")
        break
