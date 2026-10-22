import json

import docx

d = docx.Document("docs/ADELANTO DE SALARIO AGOSTO  2026.docx")

parrafos = [p.text for p in d.paragraphs if p.text.strip()]
tablas = []
for t in d.tables:
    for row in t.rows:
        tablas.append([c.text.strip() for c in row.cells])

datos = {"parrafos": parrafos, "tablas": tablas}
with open("docs/_inspect_adelanto.json", "w", encoding="utf-8") as f:
    json.dump(datos, f, ensure_ascii=False, indent=2)

print("parrafos:", len(parrafos), "| filas de tablas:", len(tablas))
