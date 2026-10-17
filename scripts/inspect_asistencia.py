import json

import openpyxl

wb = openpyxl.load_workbook("docs/Asistencia - Personal.xlsx", data_only=False)

objetivo = "Alejandra Alcaraz"
if objetivo not in wb.sheetnames:
    for n in wb.sheetnames:
        if n.strip().lower() == "alejandra alcaraz":
            objetivo = n
            break

ws = wb[objetivo]
print("Hoja:", repr(objetivo))
print("Dimensiones:", ws.dimensions, "| filas:", ws.max_row, "| cols:", ws.max_column)
print("Merged:", [str(r) for r in ws.merged_cells.ranges])
print("Orientacion:", ws.page_setup.orientation, "| paper:", ws.page_setup.paperSize)
print("Print area:", ws.print_area)
print("Col widths:", {k: v.width for k, v in ws.column_dimensions.items()})

grid = []
for r in range(1, ws.max_row + 1):
    fila = []
    for c in range(1, min(ws.max_column, 12) + 1):
        cell = ws.cell(row=r, column=c)
        fila.append(cell.value)
    grid.append(fila)

datos = {
    "hoja": objetivo,
    "dimensiones": ws.dimensions,
    "filas": ws.max_row,
    "columnas": ws.max_column,
    "merged": [str(r) for r in ws.merged_cells.ranges],
    "page": {
        "orientacion": ws.page_setup.orientation,
        "paper": ws.page_setup.paperSize,
        "margenes": {k: getattr(ws.page_margins, k, None) for k in ["left", "right", "top", "bottom"]},
    },
    "col_widths": {k: v.width for k, v in ws.column_dimensions.items()},
    "row_heights": {k: v.height for k, v in ws.row_dimensions.items() if v.height},
    "grid": grid,
}

with open("docs/_inspect_asistencia.json", "w", encoding="utf-8") as f:
    json.dump(datos, f, ensure_ascii=False, indent=2, default=str)

print("OK: docs/_inspect_asistencia.json")
