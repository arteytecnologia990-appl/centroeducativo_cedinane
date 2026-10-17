import json

import openpyxl

wb = openpyxl.load_workbook("docs/ALUMNOS.xlsx", data_only=True)
print("Hojas:", [repr(n) for n in wb.sheetnames])

objetivo = None
for n in wb.sheetnames:
    if n.strip().lower() == "cuota 2026":
        objetivo = n
        break
if not objetivo:
    for n in wb.sheetnames:
        if "cuota" in n.lower():
            objetivo = n
            break
if not objetivo:
    objetivo = wb.sheetnames[0]

ws = wb[objetivo]
print("Hoja:", repr(objetivo))
print("Dimensiones:", ws.dimensions, "| filas:", ws.max_row, "| cols:", ws.max_column)
print("Merged:", [str(r) for r in ws.merged_cells.ranges])

grid = []
for r in range(1, ws.max_row + 1):
    fila = []
    for c in range(1, min(ws.max_column, 25) + 1):
        fila.append(ws.cell(row=r, column=c).value)
    grid.append(fila)

for i, fila in enumerate(grid[:6], 1):
    print(f"  fila {i}: {fila}")

datos = {
    "hoja": objetivo,
    "dimensiones": ws.dimensions,
    "filas": ws.max_row,
    "columnas": ws.max_column,
    "merged": [str(r) for r in ws.merged_cells.ranges],
    "col_widths": {k: v.width for k, v in ws.column_dimensions.items()},
    "grid": grid,
}

with open("docs/_inspect_alumnos.json", "w", encoding="utf-8") as f:
    json.dump(datos, f, ensure_ascii=False, indent=2, default=str)

print("OK: docs/_inspect_alumnos.json")
